"""Blueprint export formats: JSON, ZIP, DOCX, PDF."""

from __future__ import annotations

import io
import json
import zipfile
from typing import Any

from app.services.export.markdown import render_blueprint


def to_json(blueprint: dict[str, Any]) -> bytes:
    return json.dumps(blueprint, indent=2, ensure_ascii=False).encode("utf-8")


def _project_slug(blueprint: dict[str, Any]) -> str:
    name = blueprint.get("project", {}).get("name", "blueprint")
    import re

    return re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")


def to_zip(blueprint: dict[str, Any]) -> bytes:
    """Bundle the blueprint as a starter repository (docs, SQL, docker assets)."""
    slug = _project_slug(blueprint)
    db = blueprint.get("database", {})
    deployment = blueprint.get("deployment", {})
    doc = blueprint.get("documentation", {})
    architecture = blueprint.get("architecture", {})
    analysis = blueprint.get("analysis", {})

    files = {
        "README.md": doc.get("readme", ""),
        "blueprint.json": json.dumps(blueprint, indent=2, ensure_ascii=False),
        "blueprint.md": render_blueprint(blueprint),
        "docs/architecture.md": doc.get("architecture_documentation", ""),
        "docs/api.md": doc.get("api_documentation", ""),
        "docs/database.md": doc.get("database_documentation", ""),
        "docs/deployment.md": doc.get("deployment_guide", ""),
        "docs/installation.md": doc.get("installation_guide", ""),
        "docs/architecture-diagrams.md": _diagrams_markdown(architecture),
        "database/schema.sql": db.get("sql_scripts", {}).get("create_tables", ""),
        "database/seed.sql": db.get("seed_data", {}).get("sql", ""),
        "database/migrations.md": "\n".join(f"- {m.get('file')}: {m.get('description')}" for m in db.get("migration_scripts", [])),
        "deploy/Dockerfile": deployment.get("dockerfile", ""),
        "deploy/docker-compose.yml": deployment.get("docker_compose", ""),
        "deploy/github-actions.yml": deployment.get("github_actions", ""),
        "deploy/env.example": _env_example(deployment),
        "docs/requirements.md": json.dumps(analysis, indent=2, ensure_ascii=False),
        "docs/domain-understanding.md": json.dumps(blueprint.get("domain_understanding", {}), indent=2, ensure_ascii=False),
        "docs/business-processes.md": json.dumps(blueprint.get("business_processes", {}), indent=2, ensure_ascii=False),
        "docs/technology-selection.md": json.dumps(blueprint.get("technology_selection", {}), indent=2, ensure_ascii=False),
        "docs/validation.md": json.dumps(blueprint.get("validation", {}), indent=2, ensure_ascii=False),
    }

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        for path, content in files.items():
            zf.writestr(f"{slug}/{path}", content if isinstance(content, str) else str(content))
    return buffer.getvalue()


def _diagrams_markdown(architecture: dict[str, Any]) -> str:
    parts = []
    for title, key in [
        ("High-Level Architecture", "high_level_architecture"),
        ("Component Diagram", "component_diagram"),
        ("Data Flow", "data_flow"),
        ("Service Communication", "service_communication"),
        ("Deployment Architecture", "deployment_architecture"),
    ]:
        mermaid = architecture.get(key, "")
        if mermaid:
            parts.append(f"## {title}\n\n```mermaid\n{mermaid}\n```\n")
    return "\n".join(parts)


def _env_example(deployment: dict[str, Any]) -> str:
    lines = ["# Generated environment template"]
    for var in deployment.get("environment_variables", []):
        lines.append(f"{var.get('name')}=")
    return "\n".join(lines) + "\n"


def to_docx(blueprint: dict[str, Any]) -> bytes:
    """Render the blueprint as a DOCX document."""
    from docx import Document

    doc = Document()
    title = blueprint.get("project", {}).get("name", "Software Blueprint")
    doc.add_heading(title, level=0)

    markdown = render_blueprint(blueprint)
    _append_markdown_to_docx(doc, markdown)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _append_markdown_to_docx(doc: Any, markdown: str) -> None:
    """Minimal markdown -> docx converter covering the blueprint renderer's output."""
    import re

    from docx.shared import Pt

    in_code = False
    for raw_line in markdown.splitlines():
        line = raw_line.rstrip()
        if line.startswith("```"):
            in_code = not in_code
            if in_code:
                doc.add_paragraph(line.replace("```", ""), style="Intense Quote")
            continue
        if in_code:
            p = doc.add_paragraph(line)
            p.paragraph_format.left_indent = Pt(18)
            continue
        if line.startswith("# "):
            doc.add_heading(line[2:], level=0)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith("#### "):
            doc.add_heading(line[5:], level=3)
        elif line.startswith("|"):
            continue
        elif re.match(r"^\- ", line) or re.match(r"^  \- ", line):
            p = doc.add_paragraph(line.lstrip(" -"))
            p.paragraph_format.left_indent = Pt(18 if line.startswith("  ") else 9)
        elif line.strip():
            doc.add_paragraph(line)


def to_pdf(blueprint: dict[str, Any]) -> bytes:
    """Render the blueprint as a PDF document."""
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        Paragraph,
        Preformatted,
        SimpleDocTemplate,
        Spacer,
    )

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, leftMargin=0.7 * inch, rightMargin=0.7 * inch)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="BlueprintCode", fontName="Courier", fontSize=7.5, leading=9.5, backColor="#F1F5F9", borderPadding=4))
    styles.add(ParagraphStyle(name="BodyTight", parent=styles["BodyText"], fontSize=9, leading=12))

    story: list[Any] = []
    title = blueprint.get("project", {}).get("name", "Software Blueprint")
    story.append(Paragraph(title, styles["Title"]))
    story.append(Spacer(1, 8))

    markdown = render_blueprint(blueprint)
    in_code = False
    code_buffer: list[str] = []
    for line in markdown.splitlines():
        if line.startswith("```"):
            in_code = not in_code
            if not in_code and code_buffer:
                story.append(Preformatted("\n".join(code_buffer), styles["BlueprintCode"]))
                story.append(Spacer(1, 6))
                code_buffer = []
            continue
        if in_code:
            code_buffer.append(line)
            continue
        if not line.strip():
            story.append(Spacer(1, 4))
            continue
        if line.startswith("### "):
            story.append(Paragraph(line[4:], styles["Heading3"]))
        elif line.startswith("## "):
            story.append(Paragraph(line[3:], styles["Heading2"]))
        elif line.startswith("# "):
            story.append(Paragraph(line[2:], styles["Heading1"]))
        elif line.startswith("- "):
            story.append(Paragraph(f"• {line[2:]}", styles["BodyTight"]))
        else:
            story.append(Paragraph(line, styles["BodyTight"]))
    if code_buffer:
        story.append(Preformatted("\n".join(code_buffer), styles["Code"]))

    doc.build(story)
    return buffer.getvalue()
